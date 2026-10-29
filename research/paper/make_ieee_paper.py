#!/usr/bin/env python3
"""Render the TRUSTXAI manuscript as an IEEE-style A4 DOCX."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "TRUSTXAI_IEEE_MANUSCRIPT.md"
OUTPUT = ROOT / "TRUSTXAI_IEEE_PAPER_SUPERVISOR_DRAFT.docx"

TITLE = "TRUSTXAI: An Evidence-Bound Framework for Auditable Local LLM Diagnostic Decision Support"
AUTHORS = [
    ("Spandan Kumar Gope", [
        "Integrated M.Tech CSE", "Computational and Data Science",
        "[School / Department, Institution]", "[City, Country]", "[Email / ORCID]",
    ]),
    ("[Supervisor Name]", [
        "[Academic designation]", "[Department]", "[Institution]",
        "[City, Country]", "[Email / ORCID]",
    ]),
    ("[Reviewer 1 Name]", [
        "[Academic designation]", "[Department]", "[Institution]",
        "[City, Country]", "[Email / ORCID]",
    ]),
    ("[Reviewer 2 Name]", [
        "[Academic designation]", "[Department]", "[Institution]",
        "[City, Country]", "[Email / ORCID]",
    ]),
]

FIGURES = {
    "<!-- FIGURE_1 -->": (
        ROOT / "figures/figure_1_architecture.png",
        "Fig. 1. TRUSTXAI local evidence and fail-closed release workflow.",
    ),
    "<!-- FIGURE_2 -->": (
        ROOT / "figures/figure_2_results.png",
        "Fig. 2. Descriptive computational outcomes. Split accuracies are not competing baselines; checker labels reflect construction intent rather than independent semantic truth.",
    ),
}

BLUE = "1F4E79"
LIGHT_BLUE = "DCEAF7"
GRAY = RGBColor(89, 89, 89)


def set_cell_margins(cell, top=40, start=55, bottom=40, end=55):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def clear_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "nil")


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_columns(section, count=2, space_twips=360):
    sect_pr = section._sectPr
    cols = sect_pr.xpath("./w:cols")
    node = cols[0] if cols else OxmlElement("w:cols")
    if not cols:
        sect_pr.append(node)
    node.set(qn("w:num"), str(count))
    node.set(qn("w:space"), str(space_twips))
    node.set(qn("w:equalWidth"), "1")


def set_font(run, size=10, bold=None, italic=None, name="Times New Roman", color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


INLINE = re.compile(r"(\*\*.+?\*\*|\*.+?\*|`.+?`)")


def add_inline(paragraph, text, size=10, base_italic=False, base_bold=False):
    text = text.replace("--", "–")
    position = 0
    for match in INLINE.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position:match.start()])
            set_font(run, size=size, italic=base_italic, bold=base_bold)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_font(run, size=size, bold=True, italic=base_italic)
        elif token.startswith("*"):
            run = paragraph.add_run(token[1:-1])
            set_font(run, size=size, italic=True, bold=base_bold)
        else:
            run = paragraph.add_run(token[1:-1])
            set_font(run, size=max(size - 0.5, 6), name="Courier New",
                     italic=base_italic, bold=base_bold)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        set_font(run, size=size, italic=base_italic, bold=base_bold)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(0.63)
    section.right_margin = Inches(0.63)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.4)
    set_columns(section, 1)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(10)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Inches(0.14)
    normal.paragraph_format.space_after = Pt(2.5)
    normal.paragraph_format.line_spacing = 1.0
    normal.paragraph_format.widow_control = True

    for style_name in ("Title", "Heading 1", "Heading 2"):
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.color.rgb = RGBColor(0, 0, 0)
    styles["Title"].font.size = Pt(20)
    styles["Title"].font.bold = False
    styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    styles["Title"].paragraph_format.space_after = Pt(8)
    title_ppr = styles["Title"]._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    styles["Heading 1"].font.size = Pt(10)
    styles["Heading 1"].font.bold = False
    styles["Heading 1"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    styles["Heading 1"].paragraph_format.space_before = Pt(8)
    styles["Heading 1"].paragraph_format.space_after = Pt(4)
    styles["Heading 1"].paragraph_format.keep_with_next = True

    styles["Heading 2"].font.size = Pt(10)
    styles["Heading 2"].font.bold = False
    styles["Heading 2"].font.italic = True
    styles["Heading 2"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    styles["Heading 2"].paragraph_format.space_before = Pt(6)
    styles["Heading 2"].paragraph_format.space_after = Pt(2)
    styles["Heading 2"].paragraph_format.keep_with_next = True

    settings = doc.settings._element
    compat = settings.find(qn("w:compat"))
    if compat is None:
        compat = OxmlElement("w:compat")
        settings.append(compat)
    do_not_expand = OxmlElement("w:doNotExpandShiftReturn")
    compat.append(do_not_expand)
    auto_hyphen = OxmlElement("w:autoHyphenation")
    auto_hyphen.set(qn("w:val"), "true")
    settings.append(auto_hyphen)


def add_author_block(doc):
    table = doc.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    clear_table_borders(table)
    table.autofit = False
    clear_table_borders(table)
    for index, (name, lines) in enumerate(AUTHORS):
        cell = table.cell(index // 2, index % 2)
        cell.width = Inches(3.42)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        set_cell_margins(cell, top=20, start=80, bottom=45, end=80)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(1)
        run = paragraph.add_run(name)
        set_font(run, size=10, bold=True)
        for line in lines:
            p = cell.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 0.92
            run = p.add_run(line)
            set_font(run, size=8.2, color=GRAY if "[" in line else None)


def extract_sections(text):
    lines = text.splitlines()
    abstract_index = lines.index("## Abstract")
    intro_index = lines.index("## I. Introduction")
    abstract_lines = []
    keywords = ""
    for line in lines[abstract_index + 1:intro_index]:
        if line.startswith("**Keywords:**"):
            keywords = line.split(":", 1)[1].strip().removeprefix("**").removesuffix("**").strip()
        elif line.strip() and not line.startswith("**"):
            abstract_lines.append(line.strip())
    return " ".join(abstract_lines), keywords, lines[intro_index:]


def add_abstract(doc, abstract, keywords):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Inches(0.15)
    paragraph.paragraph_format.space_after = Pt(7)
    label = paragraph.add_run("Abstract—")
    set_font(label, size=9, bold=True, italic=True)
    add_inline(paragraph, abstract, size=9)

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Inches(0.15)
    paragraph.paragraph_format.space_after = Pt(7)
    label = paragraph.add_run("Keywords—")
    set_font(label, size=9, bold=True, italic=True)
    add_inline(paragraph, keywords, size=9, base_italic=True)


def add_body_paragraph(doc, text, reference=False):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if reference:
        paragraph.paragraph_format.first_line_indent = Inches(-0.18)
        paragraph.paragraph_format.left_indent = Inches(0.18)
        paragraph.paragraph_format.space_after = Pt(2.5)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        paragraph.paragraph_format.line_spacing = Pt(9)
        add_inline(paragraph, text, size=8)
    else:
        paragraph.paragraph_format.first_line_indent = Inches(0.14)
        paragraph.paragraph_format.space_after = Pt(2.5)
        add_inline(paragraph, text, size=10)
    return paragraph


def add_figure(doc, marker):
    path, caption = FIGURES[marker]
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(path), width=Inches(3.22))
    caption_p = doc.add_paragraph()
    caption_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    caption_p.paragraph_format.first_line_indent = Inches(0)
    caption_p.paragraph_format.space_after = Pt(4)
    add_inline(caption_p, caption, size=8)


def table_widths(column_count):
    if column_count == 3:
        return [0.72, 1.34, 1.12]
    if column_count == 6:
        return [0.64, 0.58, 0.51, 0.54, 0.51, 0.42]
    return [3.18 / column_count] * column_count


def add_table(doc, rows):
    column_count = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=column_count)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    widths = table_widths(column_count)
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    total_twips = int(sum(widths) * 1440)
    tbl_width = tbl_pr.find(qn("w:tblW"))
    if tbl_width is None:
        tbl_width = OxmlElement("w:tblW")
        tbl_pr.append(tbl_width)
    tbl_width.set(qn("w:w"), str(total_twips))
    tbl_width.set(qn("w:type"), "dxa")
    tbl_indent = tbl_pr.find(qn("w:tblInd"))
    if tbl_indent is None:
        tbl_indent = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_indent)
    tbl_indent.set(qn("w:w"), "0")
    tbl_indent.set(qn("w:type"), "dxa")
    for col_index, width in enumerate(widths):
        table.columns[col_index].width = Inches(width)
        grid_col = table._tbl.tblGrid.gridCol_lst[col_index]
        grid_col.set(qn("w:w"), str(int(width * 1440)))
    for row_index, values in enumerate(rows):
        row = table.rows[row_index]
        prevent_row_split(row)
        if row_index == 0:
            repeat_header(row)
        for col_index, value in enumerate(values):
            cell = row.cells[col_index]
            cell.width = Inches(widths[col_index])
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_width = tc_pr.find(qn("w:tcW"))
            if tc_width is None:
                tc_width = OxmlElement("w:tcW")
                tc_pr.append(tc_width)
            tc_width.set(qn("w:w"), str(int(widths[col_index] * 1440)))
            tc_width.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, top=30, start=35, bottom=30, end=35)
            if row_index == 0:
                shade_cell(cell, LIGHT_BLUE)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if row_index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.first_line_indent = Inches(0)
            paragraph.paragraph_format.space_after = Pt(0)
            add_inline(paragraph, value, size=6.5 if column_count >= 6 else 7,
                       base_bold=row_index == 0)
    table.rows[0]._tr.get_or_add_trPr()
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_manuscript_body(doc, lines):
    index = 0
    references = False
    while index < len(lines):
        stripped = lines[index].strip()
        if not stripped:
            index += 1
            continue
        if stripped in FIGURES:
            add_figure(doc, stripped)
            index += 1
            continue
        if stripped.startswith("## "):
            heading = stripped[3:].strip()
            references = heading == "References"
            paragraph = doc.add_paragraph(style="Heading 1")
            paragraph.paragraph_format.first_line_indent = Inches(0)
            run = paragraph.add_run(heading.upper() if references else heading)
            set_font(run, size=10)
            run.font.small_caps = True
            index += 1
            continue
        if stripped.startswith("### "):
            paragraph = doc.add_paragraph(style="Heading 2")
            paragraph.paragraph_format.first_line_indent = Inches(0)
            add_inline(paragraph, stripped[4:].strip(), size=10, base_italic=True)
            index += 1
            continue
        if stripped.startswith("**TABLE ") and stripped.endswith("**"):
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Inches(0)
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.keep_with_next = True
            add_inline(paragraph, stripped[2:-2], size=8, base_bold=True)
            index += 1
            continue
        if stripped.startswith("|"):
            raw_rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                parts = [part.strip() for part in lines[index].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", part) for part in parts):
                    raw_rows.append(parts)
                index += 1
            add_table(doc, raw_rows)
            continue

        block = [stripped]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if (not candidate or candidate.startswith("## ") or candidate.startswith("### ")
                    or candidate.startswith("|") or candidate in FIGURES
                    or candidate.startswith("**TABLE ")):
                break
            block.append(candidate)
            index += 1
        add_body_paragraph(doc, " ".join(block), reference=references)


def remove_empty_first_paragraph(doc):
    if doc.paragraphs and not doc.paragraphs[0].text:
        element = doc.paragraphs[0]._element
        element.getparent().remove(element)


def main():
    manuscript = SOURCE.read_text(encoding="utf-8")
    abstract, keywords, body = extract_sections(manuscript)
    doc = Document()
    configure_document(doc)
    remove_empty_first_paragraph(doc)

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.keep_with_next = True
    run = title.add_run(TITLE)
    set_font(run, size=20)
    add_author_block(doc)
    add_abstract(doc, abstract, keywords)

    body_section = doc.add_section(WD_SECTION.CONTINUOUS)
    body_section.page_width = Cm(21.0)
    body_section.page_height = Cm(29.7)
    body_section.top_margin = Inches(0.75)
    body_section.bottom_margin = Inches(1.0)
    body_section.left_margin = Inches(0.63)
    body_section.right_margin = Inches(0.63)
    body_section.header_distance = Inches(0.3)
    body_section.footer_distance = Inches(0.4)
    set_columns(body_section, 2, 360)

    add_manuscript_body(doc, body)

    for section in doc.sections:
        section.header.is_linked_to_previous = False
        section.footer.is_linked_to_previous = False
        for paragraph in list(section.header.paragraphs) + list(section.footer.paragraphs):
            paragraph.clear()

    properties = doc.core_properties
    properties.title = TITLE
    properties.subject = "Computational feasibility study of an evidence-bound local LLM diagnostic decision-support prototype"
    properties.author = "Spandan Kumar Gope and coauthors (details pending confirmation)"
    properties.keywords = keywords
    properties.comments = "Supervisor-review draft; author affiliations and submission venue pending."

    doc.save(OUTPUT)
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    main()
